from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(r"C:\Users\HUAWEI\Downloads\DropCity_Chapter3_Methodology.docx")
OUT = Path(r"C:\Users\HUAWEI\projects\courier\DropCity_Chapter4_SystemDesignAnalysis.docx")

TRANSIT = "008080"
SLATE = "2F4F4F"
LIGHT = "F8F9FA"
BORDER = "94A3B8"


def clear_doc(doc):
    body = doc._body._element
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None, size=8.7):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color=BORDER):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_col_widths(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx >= len(row.cells):
                continue
            row.cells[idx].width = Inches(width)
            tcPr = row.cells[idx]._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(int(width * 1440)))
            tcW.set(qn("w:type"), "dxa")


def para(doc, text="", style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.name = "Arial"
    return p


def heading(doc, text, level):
    return doc.add_paragraph(text, style=f"Heading {level}")


def caption(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    run = p.runs[0]
    run.font.name = "Arial"
    run.font.size = Pt(9.5)
    run.italic = True


def table_caption(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.runs[0]
    run.font.name = "Arial"
    run.font.size = Pt(9.5)
    run.bold = True


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for idx, header in enumerate(headers):
        shade_cell(t.rows[0].cells[idx], SLATE)
        set_cell_text(t.rows[0].cells[idx], header, bold=True, color="FFFFFF", size=9)
    for row in rows:
        cells = t.add_row().cells
        for idx, value in enumerate(row):
            shade_cell(cells[idx], LIGHT if len(t.rows) % 2 == 0 else "FFFFFF")
            set_cell_text(cells[idx], value)
    set_col_widths(t, widths)
    set_table_borders(t)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return t


def placeholder(doc, title, description, fig_no, cap, analysis):
    t = doc.add_table(rows=2, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_col_widths(t, [6.3])
    shade_cell(t.rows[0].cells[0], SLATE)
    set_cell_text(t.rows[0].cells[0], f"DIAGRAM / SCREENSHOT PLACEHOLDER: {title}", bold=True, color="FFFFFF", size=9.4)
    shade_cell(t.rows[1].cells[0], LIGHT)
    set_cell_text(t.rows[1].cells[0], f"Visual Description: {description}", size=8.8)
    set_table_borders(t)
    caption(doc, f"Figure 4.{fig_no}: {cap}")
    para(doc, analysis)


def bullets(doc, items):
    for item in items:
        try:
            p = doc.add_paragraph(style="List Bullet")
        except Exception:
            p = doc.add_paragraph(style="List Paragraph")
        p.add_run(item).font.name = "Arial"
        p.paragraph_format.space_after = Pt(4)


doc = Document(str(BASE))
clear_doc(doc)
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)

for style_name in ["Normal", "List Paragraph"]:
    try:
        st = doc.styles[style_name]
        st.font.name = "Arial"
        st.font.size = Pt(11)
        st.paragraph_format.line_spacing = 1.08
        st.paragraph_format.space_after = Pt(8)
    except Exception:
        pass

heading(doc, "CHAPTER 4: SYSTEM DESIGN AND ANALYSIS", 1)
heading(doc, "4.1  Introduction", 2)
para(doc, "This chapter converts the methodology and technical architecture presented in Chapter 3 into a concrete systems analysis and design specification for DropCity. The purpose is not to repeat the algorithmic proof, but to show how the researched ideas are organised into deployable components: Flutter mobile nodes, an Express API gateway, PostGIS persistence, Firebase-backed notification services, WebSocket update channels, and background workers monitored through a heartbeat watchdog.")
para(doc, "The analysis follows the DCE project template structure for analysis and design: hardware and software requirements, user and domain requirements, component identification, architectural design, database design, interface design, security design, and implementation-readiness checks. Each technical claim in this chapter was reconciled against the current repository implementation, especially the backend SQL migrations, route modules, service layer, scheduler, and courier tracking service.")
para(doc, "A small but important distinction is maintained throughout the chapter. Some state names are used as dissertation-level analytical abstractions, while the running code persists more compact values in PostgreSQL. For example, the tracking lifecycle is analysed as IDLE -> CORRIDOR_TRACKING -> TRAFFIC_DETOUR -> DEVIATION_ALERT -> ARRIVED, while the implemented client-visible tracking summary currently stores values such as NOMINAL, ON_CORRIDOR, MOVING_POSITIVELY, and OFF_CORRIDOR_STATIONARY.")

heading(doc, "4.2  Hardware and Software Specifications", 2)
para(doc, "DropCity is a distributed system rather than a single desktop application. The specification therefore separates the development workstation, mobile execution nodes, server runtime, database tier, and administrator console. Hardware entries are presented as design requirements rather than measured performance results, because final production device qualification must be completed during field deployment.")
table_caption(doc, "Table 4.1: Development and administration computer specifications")
table(doc, ["Component", "Specification", "Design Rationale"], [
    ["Processor", "Modern 64-bit CPU capable of running Node.js, Flutter tooling, and browser-based admin workloads", "The project requires simultaneous backend execution, mobile builds, and dashboard testing."],
    ["Memory", "Sufficient RAM for Android tooling, backend process, and browser console to run concurrently", "Avoids build-time instability during Flutter and Next.js validation."],
    ["Storage", "SSD-backed workspace with room for source code, dependencies, emulators, logs, and generated documentation", "Node modules, Flutter build artifacts, and rendered document assets grow quickly."],
    ["Network", "Stable internet access for Firebase, Supabase, package tooling, and map-provider requests", "The local machine coordinates cloud-backed services during development and demonstration."],
    ["Display", "Readable desktop display for admin console, emulator, and documentation review", "Supports side-by-side verification of mobile, web, and backend evidence."],
], [1.35, 2.55, 2.25])

table_caption(doc, "Table 4.2: Courier and client mobile device requirements")
table(doc, ["Requirement", "Courier App", "Client App"], [
    ["Operating environment", "Android/iOS device supported by Flutter and background geolocation dependencies", "Android/iOS device supported by Flutter client dependencies"],
    ["Location capability", "GPS-capable device with foreground and background location permission", "GPS-capable device for pickup/dropoff selection and status context"],
    ["Camera", "Required for pickup/dropoff photographic evidence", "Required when sender or recipient flow needs proof capture"],
    ["Connectivity", "Mobile data or Wi-Fi, with local SQLite outbox fallback during disconnection", "Mobile data or Wi-Fi for notifications, status refresh, and handoff flows"],
    ["Local persistence", "SQLite outbox for tracking pulses and dead-letter queue", "Shared preferences and app state for session and UI continuity"],
], [1.55, 2.35, 2.25])
para(doc, "The courier device is the more demanding edge node because it performs continuous geolocation, local queueing, connectivity checks, and retry scheduling. The client device performs lighter delivery-creation, status-viewing, handoff, and notification operations.")

table_caption(doc, "Table 4.3: Verified software stack and runtime dependencies")
table(doc, ["Layer", "Verified Technology", "Repository Evidence"], [
    ["Backend API", "Node.js with Express 4.19.2", "backend package manifest defines Express API entrypoint at src/index.js."],
    ["Database", "Supabase PostgreSQL with PostGIS SQL migrations", "backend/sql contains corridor, parcel, tracking, route, heartbeat, and connectivity migrations."],
    ["Mobile clients", "Flutter SDK constraint >=3.3.0 <4.0.0", "client and courier pubspec files define the Flutter app boundaries."],
    ["Courier tracking", "flutter_background_geolocation 5.1.2, geolocator 13.0.2, sqflite 2.3.3+1", "Courier runtime includes background geolocation and local tracking_outbox tables."],
    ["Admin console", "Next.js 14.2.35, React 18.3.1, TypeScript", "admin package manifest defines the operations dashboard stack."],
    ["Notifications", "Firebase Admin SDK 12.5.0 and Firebase Messaging packages", "Backend and mobile manifests include Firebase services for topic and device-token flows."],
    ["Background jobs", "node-cron 3.0.3", "Scheduler registers matching, tracking, cleanup, scoring, notification, and watchdog jobs."],
    ["Real-time channel", "ws 8.18.0 WebSocket server", "Backend broadcasts parcel status, handshake events, and tracking updates."],
], [1.35, 2.45, 2.35])

table_caption(doc, "Table 4.4: Browser and administration interface requirements")
table(doc, ["Access Role", "Interface", "Requirement"], [
    ["Administrator", "Next.js admin console", "Modern Chromium, Firefox, or Edge browser with JavaScript enabled."],
    ["Operations reviewer", "Admin monitoring, alerts, disputes, health, and spatial analytics pages", "Authenticated Supabase-backed session and network access to backend API."],
    ["Developer", "Local browser and emulator", "Browser developer tools for API, WebSocket, and UI state inspection."],
], [1.45, 2.25, 2.45])

heading(doc, "4.3  Problem Domain and User Requirements", 2)
para(doc, "The design problem is a constrained urban logistics problem: parcels must move across Harare without creating dedicated delivery trips, without exposing a courier's raw live location to clients, and without losing operational integrity when the mobile network is unreliable. The system must therefore support route declaration, corridor-constrained matching, physical handoff verification, background tracking, offline recovery, and administrative observability.")
heading(doc, "4.3.1  Functional Requirements", 3)
table_caption(doc, "Table 4.5: Functional requirements mapped to implemented modules")
table(doc, ["Code", "Functional Requirement", "Confirmed Implementation Surface"], [
    ["FR1", "Allow clients to create parcel delivery requests with origin, destination, parcel attributes, and recipient context.", "POST /parcels and client delivery creation flow."],
    ["FR2", "Allow couriers to declare, activate, and complete planned routes.", "POST /couriers/routes, PATCH /couriers/routes/:routeId, route status PLANNED/ACTIVE/COMPLETED/CANCELLED."],
    ["FR3", "Match parcels only to feasible courier corridors and rank assignment candidates.", "match_delivery_to_couriers SQL RPC and backend matching service filtering active/about-to-start routes."],
    ["FR4", "Verify pickup using PIN, proximity gate, and photo evidence before tracking begins.", "POST /handshake/pickup updates parcel status to IN_TRANSIT after checks pass."],
    ["FR5", "Verify dropoff using PIN or recipient/manual OTP, proximity gate, and photo evidence.", "POST /handshake/dropoff and POST /handshake/courier/complete-dropoff update status to COMPLETED."],
    ["FR6", "Accept live tracking pulses and generate privacy-preserving progress summaries.", "POST /tracking/update writes courier_tracking_logs and updates parcels tracking summary fields."],
    ["FR7", "Recover tracking pulses captured during network loss.", "Courier SQLite tracking_outbox flushes to POST /tracking/batch-sync in FIFO order."],
    ["FR8", "Expose operational health, alerts, disputes, and tracking observability to administrators.", "Admin route group includes health heartbeats, disputes, alerts, jobs, tracking observability, and spatial analytics."],
], [0.7, 3.0, 2.45])

heading(doc, "4.3.2  Non-Functional Requirements", 3)
table_caption(doc, "Table 4.6: Non-functional requirements and design responses")
table(doc, ["Quality Attribute", "Requirement", "Design Response"], [
    ["Reliability", "Background jobs must be monitored for missed execution.", "job_heartbeats table and heartbeat_watchdog job mark stale workers as STUCK after 2x expected frequency."],
    ["Privacy", "Clients should not receive raw courier coordinates by default.", "Client-facing status uses progress percentage and integrity labels instead of raw location streams."],
    ["Resilience", "Tracking must survive temporary disconnection.", "Courier app stores unsent pulses in SQLite and retries with capped exponential backoff."],
    ["Security", "Custody transfer must resist casual spoofing.", "PIN hashes, GPS proximity gates, role checks, and photo evidence combine into the 3WH evidence chain."],
    ["Scalability", "Spatial matching must avoid full manual route inspection.", "PostGIS geography columns, GiST indexes, ST_DWithin, and ST_LineLocatePoint reduce matching to spatial filtering."],
    ["Observability", "Administrators must detect stuck jobs and delivery anomalies.", "Admin health, alerts, logs, disputes, and tracking observability routes expose operational state."],
], [1.35, 2.35, 2.45])
placeholder(doc, "DropCity Use-Case Boundary", "A UML use-case diagram with four actors: Client, Courier, Recipient, and Administrator. Client use cases: create delivery, view courier acceptance, track progress, issue rating, report handoff issue. Courier use cases: declare route, accept parcel, complete pickup, stream tracking pulses, complete dropoff, view alerts. Recipient use cases: issue dropoff OTP and confirm handoff. Administrator use cases: approve couriers/vehicles, monitor heartbeats, investigate disputes, resolve alerts, inspect spatial analytics.", 1, "Use-case boundary for the DropCity distributed logistics platform, showing how each actor interacts with the implemented API modules.", "The use-case model clarifies that DropCity is not merely a client-courier marketplace. It is a four-party custody and observability system in which recipient and administrator actions are first-class requirements, especially where physical handoff evidence and operational reliability are concerned.")

heading(doc, "4.4  System Components and Functional Decomposition", 2)
para(doc, "The implemented repository is organised around four deployable surfaces: the client Flutter app, the courier Flutter app, the Next.js admin console, and the Express backend. The backend then decomposes into route modules, service modules, SQL migrations, background jobs, WebSocket broadcasting, and cloud integration utilities.")
table_caption(doc, "Table 4.7: Component decomposition of the implemented system")
table(doc, ["Component", "Primary Responsibility", "Key Implementation Evidence"], [
    ["Client Flutter app", "Delivery creation, status viewing, sender/recipient handoff screens, rating, notifications", "client/lib/screens includes delivery_creation_flow_screen, parcel_status_screen, recipient_handoff_screen, and sender fallback screens."],
    ["Courier Flutter app", "Route declaration, assigned parcel workflow, pickup/dropoff, background tracking, local outbox", "courier/lib/screens and courier/lib/services/courier_tracking_service.dart."],
    ["Admin console", "Monitoring, logs, alerts, health, disputes, vehicles, spatial analytics", "admin/src/app pages and backend /admin route group."],
    ["Express API gateway", "Authentication boundary, mounted route modules, request IDs, JSON API responses", "backend/src/index.js mounts /auth, /users, /vehicles, /corridors, /parcels, /matches, /handshake, /tracking, /admin, and related modules."],
    ["PostGIS persistence", "Spatial routes, parcel points, virtual handoff points, tracking logs, heartbeat states", "backend/sql migrations define corridors, parcels, courier_tracking_logs, routes, route_deviation_events, and job_heartbeats."],
    ["Background scheduler", "Matching, alerts, tracking validation, notification outbox, cleanup, scoring, watchdog", "backend/src/services/scheduler.js registers recurring jobs through node-cron."],
    ["WebSocket channel", "Low-latency parcel status, handshake, and tracking-update broadcasts", "backend/src/ws.js broadcasts parcel_status, handshake_event, and tracking_update payloads."],
], [1.45, 2.5, 2.2])
placeholder(doc, "Component Interaction Diagram", "A layered component diagram. Top layer: Client Flutter Node, Courier Flutter Node, and Next.js Admin Console. Middle layer: Node.js Express API Gateway with route modules. Side channel: WebSocket broadcaster connected to Client and Courier. Service layer: matching service, notification service, ETA model, heuristic tracking, scheduler. Data layer: Supabase/PostgreSQL/PostGIS, Firebase messaging, and local SQLite outbox on courier device.", 2, "Logical component interaction diagram showing how mobile, web, API, service, and persistence components collaborate.", "The diagram should make visible the central architectural decision: security-critical and spatially authoritative operations are performed server-side, while mobile devices behave as edge nodes that collect inputs, cache telemetry, and display privacy-filtered summaries.")

heading(doc, "4.5  System Architecture and Design Considerations", 2)
heading(doc, "4.5.1  Context and Data Flow Design", 3)
para(doc, "At context level, DropCity receives delivery intent from the client, route intent from the courier, and operational supervision from the administrator. The API gateway validates identity and role, converts user actions into database mutations, and broadcasts summary events to subscribed devices. The data flow deliberately separates private operational evidence from public delivery status: raw geospatial pulses remain in courier_tracking_logs, while progress_percent and integrity_status are exposed to the client-facing parcel status view.")
placeholder(doc, "Level-0 Data Flow Diagram", "A DFD with external entities Client, Courier, Recipient, Administrator, Firebase Messaging, and Supabase/PostGIS. Processes: Create Delivery, Declare Route, Match Corridor, Verify Pickup, Track In Transit, Verify Dropoff, Monitor Operations. Data stores: parcels, corridors, parcel_assignment_queue, courier_tracking_logs, route_deviation_events, job_heartbeats, notification_outbox. Show raw location flowing only into private tracking logs and summary status flowing back to clients.", 3, "Level-0 data flow diagram for DropCity, emphasizing privacy separation between raw telemetry and client-visible progress summaries.", "This DFD is central to the dissertation argument because it shows how the platform balances transparency and privacy. The client receives enough information to trust delivery progress, but not enough raw geospatial detail to continuously surveil the courier.")

heading(doc, "4.5.2  API Gateway and Route Mounting", 3)
para(doc, "The Express entrypoint mounts each controller under a domain-oriented prefix. Public health and authentication routes are available before the user middleware, while operational modules run behind authentication and, where necessary, role checks inside the route file. Request IDs are attached to JSON responses so errors can be traced across mobile clients, backend logs, and administrative evidence screens.")
table_caption(doc, "Table 4.8: Implemented backend API route groups")
table(doc, ["Route Group", "Primary Use", "Examples of Implemented Endpoints"], [
    ["/auth", "Authentication and session flows", "Signup/login/password flows through backend auth module."],
    ["/users and /vehicles", "Profiles, role setup, vehicle metadata", "Courier onboarding and admin vehicle verification."],
    ["/corridors and /couriers", "Route geometry and courier lifecycle", "Route declaration, route templates, route activation and completion."],
    ["/parcels", "Parcel lifecycle and assignment", "Create parcel, assign/request courier, accept/decline, privacy, checkpoints, rating."],
    ["/matches", "Spatial matching", "POST /matches/corridors and POST /matches/delivery."],
    ["/handshake", "Pickup/dropoff custody gates", "POST /handshake/init, /pickup, /dropoff, recipient/manual OTP completion."],
    ["/tracking", "Live and offline tracking", "POST /tracking/update, POST /tracking/batch-sync, /alerts/me, /heuristics/:parcelId, /eta/:parcelId."],
    ["/heartbeat and /admin", "Reliability and operations", "POST /heartbeat, /admin/health/heartbeats, /admin/tracking/observability, /admin/disputes."],
], [1.3, 2.25, 2.6])

heading(doc, "4.5.3  Physical Deployment Design", 3)
para(doc, "The physical design separates user devices, API execution, and managed persistence. Flutter apps operate on mobile devices and communicate with the backend over authenticated HTTP and WebSocket channels. The backend runs as a Node.js process that owns the business rules, background jobs, and integration with Supabase and Firebase. Supabase/PostGIS provides authoritative geospatial and transactional storage, while SQLite exists only as a courier-side resilience buffer during temporary connectivity loss.")
placeholder(doc, "Physical Deployment Topology", "A deployment diagram with three zones. Edge zone: Client Phone and Courier Phone; Courier Phone contains SQLite tracking_outbox and background geolocation service. Application zone: Express API Gateway, WebSocket server, background scheduler, notification outbox processor. Data/cloud zone: Supabase PostgreSQL/PostGIS, Firebase Cloud Messaging, Admin Console hosting. Mark HTTPS API calls, WebSocket subscriptions, and FIFO batch-sync arrows.", 4, "Physical deployment topology for the DropCity mobile-edge, API, and cloud-persistence architecture.", "The physical topology explains why DropCity is classified as a distributed system. A courier phone can continue collecting ordered telemetry while disconnected, then reconcile with the central platform once a usable network path returns.")

heading(doc, "4.6  Database and Spatial Design", 2)
para(doc, "The database design is built around spatial authority. Corridors, parcel endpoints, pickup/dropoff points, and tracking pulses are stored using geography point or line objects, and spatial indexes are applied to key route and parcel columns. This allows the matching function and tracking validator to reason about location as first-class data rather than as ordinary text coordinates.")
table_caption(doc, "Table 4.9: Core database entities and design purpose")
table(doc, ["Entity", "Purpose", "Important Fields or Constraints"], [
    ["job_heartbeats", "Stores backend worker health state", "job_name primary key, expected_frequency_sec, last_heartbeat_at, status constrained to ACTIVE/STUCK/INACTIVE_CLEAN."],
    ["corridors", "Stores declared route geometry", "start_point, end_point, corridor_line, window_start/window_end, GiST indexes on spatial columns."],
    ["parcels", "Stores delivery lifecycle and handoff evidence", "origin/destination points, pickup/dropoff points, PIN hashes, photo URLs, verification timestamps, privacy_mode, assigned_courier_id, tracking summary fields."],
    ["parcel_assignment_queue", "Stores ranked parcel-corridor assignment candidates", "parcel_id, corridor_id, rank, status."],
    ["handshake_events", "Stores auditable custody-transfer events", "parcel_id, step, actor_id, status, lat/lng/accuracy, photo_url, created_at."],
    ["courier_tracking_logs", "Stores private raw tracking pulses", "raw_location geography point, is_on_corridor, progress_index, distance metrics, device_info, network_info, heuristic_flags, speed_kmh."],
    ["route_deviation_events", "Stores route-integrity alerts", "deviation_type, duration_seconds, last_on_corridor_at, resolution fields."],
    ["connectivity_map and zone_connectivity", "Stores logistics-as-a-sensor observations", "network provider/device model/signal bucket and aggregated zone-level reliability counters."],
], [1.45, 2.35, 2.35])

heading(doc, "4.6.1  Zero-Detour Spatial Matching Design", 3)
para(doc, "The current SQL matching function receives a parcel origin, destination, and maximum detour distance. It returns corridor candidates only where both endpoints lie within the accepted spatial corridor and the pickup projection appears before the dropoff projection along the route line. In mathematical form, a candidate corridor c is valid only when:")
para(doc, "$$ST_DWithin(c, L_p, W) \\land ST_DWithin(c, L_d, W) \\land ST_LineLocatePoint(c, L_p) < ST_LineLocatePoint(c, L_d)$$")
para(doc, "The implemented backend then enriches this base spatial result by checking route lifecycle and capacity before producing ranked courier options. This prevents a geometrically valid but operationally unavailable route from becoming an assignment candidate.")
placeholder(doc, "Database Entity Relationship Diagram", "An ERD centered on parcels. Connect parcels to users through created_by, recipient_id, and assigned_courier_id. Connect parcels to parcel_assignment_queue, corridors, handshake_events, courier_tracking_logs, route_deviation_events, eta_calculations_log, and parcel_checkpoints. Show corridors connected to routes and courier users. Highlight geospatial columns with a small map-pin icon and heartbeat/job tables as operational metadata.", 5, "Database entity relationship model for DropCity's parcel, corridor, handoff, tracking, and operational-health records.", "The ERD should help an assessor see that DropCity's database is designed around custody evidence and spatial reasoning, not only around ordinary CRUD records. The route, parcel, tracking, and handoff entities together form the audit trail required for an opportunistic logistics platform.")

heading(doc, "4.7  Tracking Lifecycle State Machine", 2)
para(doc, "The dissertation-level tracking lifecycle is defined as the ordered state set $S = \\{IDLE, CORRIDOR_TRACKING, TRAFFIC_DETOUR, DEVIATION_ALERT, ARRIVED\\}$. This abstraction describes how the system should be analysed during delivery, while the repository currently persists implementation-level summary labels on the parcel record and raw pulse diagnostics in courier_tracking_logs.")
table_caption(doc, "Table 4.10: Analytical tracking state machine mapped to implementation fields")
table(doc, ["Analytical State", "Entry Condition", "Implemented Representation", "System Action"], [
    ["IDLE", "Parcel is not yet in the active transit window or no usable pulse has been received.", "Parcel status before pickup verification, or tracking_integrity_status null/NOMINAL.", "Wait for pickup verification or first valid tracking pulse."],
    ["CORRIDOR_TRACKING", "Courier pulse is within the corridor tolerance and progress index can be computed.", "tracking_integrity_status = ON_CORRIDOR; courier_tracking_logs.is_on_corridor = true.", "Update progress percentage and broadcast tracking_update summary."],
    ["TRAFFIC_DETOUR", "Courier is outside corridor but distance-to-destination is improving.", "tracking_integrity_status = MOVING_POSITIVELY.", "Preserve delivery trust while marking that route geometry no longer matches the planned corridor."],
    ["DEVIATION_ALERT", "Courier remains off-corridor/stalled beyond validation thresholds.", "tracking_integrity_status = OFF_CORRIDOR_STATIONARY and/or route_deviation_events entry.", "Notify courier/admin-facing channels and expose alert history."],
    ["ARRIVED", "Dropoff gate succeeds and custody is completed.", "Parcel status = COMPLETED; dropoff_verified_at populated.", "Stop active tracking relevance and preserve evidence chain."],
], [1.25, 1.75, 1.8, 1.35])
para(doc, "The transition function $\\delta(s_t, x_t)$ is driven by four inputs: parcel handoff status, pulse availability, corridor containment, and movement diagnostics. Corridor containment uses the configured tolerance in the tracking route, while movement diagnostics are derived from destination distance, previous pulse history, speed, and heuristic flags.")
placeholder(doc, "Tracking Lifecycle State Machine", "A state machine diagram with five large states: IDLE, CORRIDOR_TRACKING, TRAFFIC_DETOUR, DEVIATION_ALERT, ARRIVED. Transitions: pickup_verified_at set -> CORRIDOR_TRACKING; is_on_corridor=false and movement positive -> TRAFFIC_DETOUR; off-corridor or no recent pulse exceeds threshold -> DEVIATION_ALERT; dropoff_verified_at set -> ARRIVED; recovery pulse on corridor -> CORRIDOR_TRACKING. Annotate implementation labels ON_CORRIDOR, MOVING_POSITIVELY, OFF_CORRIDOR_STATIONARY, COMPLETED next to the relevant states.", 6, "Tracking lifecycle state machine showing the relationship between dissertation states and implemented persistence labels.", "This state machine is important because it makes the privacy-preserving design operational. Clients do not need raw latitude and longitude to understand delivery progress; they need a trustworthy state label, a progress percentage, and freshness information.")

heading(doc, "4.7.1  Offline Tracking and ARQ Design", 3)
para(doc, "The courier app implements a local Asynchronous Request Queue for tracking pulses. When a live update fails because the device has no network or the HTTP request times out, the pulse is inserted into the SQLite tracking_outbox table with parcel_id, latitude, longitude, accuracy, timestamp, retry_count, next_attempt_at, last_error, and created_at. The flush routine reads rows in ascending id order, batches them, posts them to /tracking/batch-sync, deletes successful or permanently skipped records, schedules retries for recoverable failures, and moves exhausted records into tracking_dead_letter.")
table_caption(doc, "Table 4.11: Courier tracking outbox resilience rules")
table(doc, ["Rule", "Implemented Behaviour", "Design Benefit"], [
    ["FIFO ordering", "fetchBatch orders by id ASC.", "Preserves chronological interpretation of courier movement."],
    ["Age control", "Client removes old rows; backend rejects updates older than 24 hours.", "Prevents stale telemetry from corrupting current delivery state."],
    ["Retry control", "Failed rows receive next_attempt_at with capped exponential backoff.", "Avoids network hammering while preserving eventual delivery."],
    ["Dead-letter handling", "Rows exceeding retry limits move to tracking_dead_letter.", "Keeps the main queue healthy and exposes unresolved sync failures."],
    ["Batch sync", "Courier posts queued updates to /tracking/batch-sync.", "Reduces overhead after reconnecting from a network outage."],
], [1.25, 2.55, 2.35])

heading(doc, "4.8  Watchdog Heartbeat Logic", 2)
para(doc, "The backend reliability design treats scheduled background jobs as observable workers. Each job is expected to write a heartbeat record after execution. The watchdog reads job_heartbeats, computes the elapsed time since last_heartbeat_at, and marks workers as STUCK when the elapsed time exceeds twice the expected frequency.")
para(doc, "Formally, for worker j with expected period f_j seconds and current time t_n, the watchdog predicate is:")
para(doc, "$$STUCK(j) \\iff (t_n - last_heartbeat_at(j)) > 2f_j$$")
table_caption(doc, "Table 4.12: Watchdog heartbeat states")
table(doc, ["State", "Trigger Condition", "Persisted Action"], [
    ["ACTIVE", "Worker records a heartbeat within the expected frequency window.", "Update last_heartbeat_at and status ACTIVE."],
    ["STUCK", "Worker misses more than 2x expected_frequency_sec.", "heartbeat_watchdog updates status to STUCK and last_status_change_at."],
    ["INACTIVE_CLEAN", "Worker is intentionally taken out of service for controlled maintenance.", "Schema allows this non-error terminal state to avoid false stuck alerts."],
], [1.3, 2.7, 2.1])
table_caption(doc, "Table 4.13: Default scheduled backend jobs")
table(doc, ["Job", "Frequency", "Purpose"], [
    ["heartbeat_watchdog", "Every 5 minutes", "Monitors heartbeat freshness and marks stale jobs as STUCK."],
    ["check_alerts", "Every 2 minutes", "Evaluates alert rules and notification conditions."],
    ["match_parcels", "Every 3 minutes", "Matches pending parcels with available corridors."],
    ["validate_tracking", "Every 2 minutes", "Detects route deviations and stalled tracking for in-transit parcels."],
    ["cleanup_tracking_data", "Daily at 02:00", "Deletes old tracking/deviation records according to retention policy."],
    ["process_notification_outbox", "Every 30 seconds", "Processes notification outbox rows and retries failed deliveries."],
    ["aggregate_courier_scores", "Hourly", "Recomputes courier score from ratings, adherence, punctuality, and incidents."],
    ["heuristic_tracking_sweep", "Every 3 minutes", "Runs heuristic tracking analysis for in-transit parcels."],
], [2.0, 1.35, 2.8])
placeholder(doc, "Watchdog Heartbeat Sequence", "A sequence diagram with lifelines: Background Worker, Scheduler, job_heartbeats table, Heartbeat Watchdog, Admin Health Dashboard. Normal path: worker completes job, scheduler upserts ACTIVE heartbeat, admin dashboard reads healthy status. Failure path: worker stops emitting pulses, watchdog compares now - last_heartbeat_at against 2x expected_frequency_sec, updates status to STUCK, admin dashboard displays the stuck worker.", 7, "Watchdog sequence diagram showing how missed heartbeats transition workers from ACTIVE to STUCK in the database.", "The watchdog design converts invisible backend failure into explicit database state. This is essential for DropCity because matching, tracking validation, notifications, and cleanup are background responsibilities that may fail silently if not externally monitored.")

heading(doc, "4.9  Interface Design", 2)
para(doc, "The interface design follows the same role separation as the backend. The client app prioritises delivery creation, status readability, and recipient/sender handoff actions. The courier app prioritises route declaration, assigned parcel execution, pickup/dropoff proof capture, background tracking health, and alert visibility. The admin console prioritises operational supervision: heartbeats, logs, vehicles, alerts, disputes, courier review, spatial analytics, and scheduler state.")
table_caption(doc, "Table 4.14: Interface design matrix")
table(doc, ["Interface", "Key Screens or Pages", "Design Purpose"], [
    ["Client app", "Home, Deliveries, Account, delivery creation flow, parcel status, recipient handoff, sender fallback", "Simple delivery creation and transparent but privacy-preserving progress monitoring."],
    ["Courier app", "Home, Routes, Parcels, Settings, route declaration, assigned parcels, pickup mode, pickup, dropoff", "Fast operational workflow that minimises interaction while the courier is moving."],
    ["Admin console", "Monitoring, alerts, logs, couriers, vehicles, health, disputes, scheduler, spatial analytics", "Central control plane for reliability, compliance, disputes, and field diagnostics."],
], [1.35, 2.8, 2.0])
placeholder(doc, "Three-Panel Interface Wireframe", "A three-panel figure. Left: Client Parcel Status screen showing status badge, progress bar, ETA, tracking integrity label, and handoff timeline. Middle: Courier Pickup Mode screen with assigned parcel card, pickup action, camera proof, and route alert chip. Right: Admin Health dashboard showing job heartbeat cards, stuck job count, tracking observability, and alert history.", 8, "Interface wireframe blueprint for the client, courier, and admin surfaces.", "The interface blueprint demonstrates that each user sees only the information necessary for their role. This reduces cognitive load and strengthens privacy by preventing client screens from becoming raw surveillance dashboards.")

heading(doc, "4.10  Security Design", 2)
heading(doc, "4.10.1  Three-Way Handshake Security", 3)
para(doc, "The DropCity handoff security model combines something the actor knows, somewhere the actor physically is, and something the actor captures. The PIN/OTP is hashed before storage, the GPS coordinate must be close to the pickup or dropoff point, and the handoff must include photographic evidence. Successful pickup updates pickup_verified_at and moves the parcel to IN_TRANSIT; successful dropoff populates dropoff_verified_at and moves the parcel to COMPLETED.")
table_caption(doc, "Table 4.15: Three-Way Handshake security checks")
table(doc, ["Gate", "Verification Inputs", "Successful State Mutation"], [
    ["PIN/OTP gate", "Pickup PIN hash, dropoff PIN hash, recipient-issued OTP hash, or manual dropoff OTP hash", "Failed attempts increment counters; successful verification clears or resets relevant fields."],
    ["Spatial gate", "Courier latitude/longitude compared against pickup, meeting pickup, or dropoff point", "Rejected when the actor is outside the accepted proximity boundary."],
    ["Evidence gate", "pickup_photo_url or dropoff_photo_url", "Photo URL is stored with the parcel and mirrored in handshake_events."],
    ["Role gate", "Authenticated user role and assigned_courier_id / recipient_id checks", "Prevents unrelated users from completing pickup or dropoff actions."],
], [1.25, 3.0, 1.9])
heading(doc, "4.10.2  Privacy and Operational Security", 3)
para(doc, "Privacy is designed into both the data model and the event model. Raw courier coordinates are recorded privately for route integrity, dispute investigation, and spatial analytics. Client-facing updates are reduced to status, progress percentage, freshness, ETA, and integrity labels. Administrators can inspect richer evidence through restricted admin endpoints, but ordinary clients do not receive the courier's raw movement trail.")
bullets(doc, [
    "API requests are wrapped with request IDs for traceability across logs and JSON responses.",
    "Business routes run behind authentication middleware, with role-specific checks applied in sensitive controllers.",
    "The courier tracking endpoint rejects updates from couriers not assigned to the parcel.",
    "The heartbeat route applies conflict checks to reduce accidental job-name takeover.",
    "The notification outbox decouples event creation from delivery retries so user actions do not block on push delivery.",
])

heading(doc, "4.11  Design Verification and Test Plan", 2)
para(doc, "Chapter 5 will present implementation results and measured performance. Chapter 4 therefore limits testing content to design verification: the checks required to prove that each designed module behaves according to its specification. This avoids presenting unverified performance statistics before the final pilot evidence is gathered.")
table_caption(doc, "Table 4.16: System design verification test plan")
table(doc, ["Module", "Test Focus", "Expected Verification Evidence"], [
    ["Authentication and roles", "Client, courier, and admin requests reach only permitted route groups.", "Forbidden responses for wrong-role actions; successful responses for valid users."],
    ["Route declaration", "Courier can create, activate, and complete routes.", "Route status transitions PLANNED -> ACTIVE -> COMPLETED or CANCELLED."],
    ["Spatial matching", "Parcel endpoints must lie within corridor and pickup fraction must precede dropoff fraction.", "Matching result contains candidate corridor, projected pickup/dropoff points, and ordered fractions."],
    ["Pickup handshake", "PIN, GPS, photo, assignment, and status mutation are enforced.", "Parcel status becomes IN_TRANSIT and handshake_events records SUCCESS."],
    ["Tracking update", "Assigned courier pulse creates private log and public summary.", "courier_tracking_logs row exists; parcels tracking_progress_percent and tracking_integrity_status are updated."],
    ["Offline batch sync", "Queued pulses flush in FIFO order after connectivity returns.", "Batch response reports synced/skipped rows; local outbox deletes successful rows."],
    ["Deviation detection", "Off-corridor or stale tracking raises deviation evidence.", "route_deviation_events contains OFF_CORRIDOR_PROLONGED or NO_RECENT_TRACKING when threshold is crossed."],
    ["Watchdog", "Missed worker heartbeats become visible to admin.", "job_heartbeats status changes from ACTIVE to STUCK after the 2x threshold."],
    ["Dropoff handshake", "Dropoff OTP/PIN, GPS, and photo complete custody chain.", "Parcel status becomes COMPLETED and dropoff evidence is stored."],
], [1.45, 2.55, 2.15])

heading(doc, "4.12  Implementation Constraints and Risk Analysis", 2)
para(doc, "The current system design is strong enough for end-to-end pilot operation, but several risks must remain visible in the dissertation. Android background execution is affected by device manufacturer battery policies, meaning long-haul app-closed tracking still requires field validation on the target handset set. The repository also distinguishes between design-level Firestore/public-channel language and the confirmed implementation, where parcel tracking summaries are stored in Supabase and broadcast through WebSockets. This distinction should be preserved during assessment to avoid overstating the live implementation.")
table_caption(doc, "Table 4.17: Implementation risks and mitigation strategies")
table(doc, ["Risk", "Likely Impact", "Mitigation in Current Design"], [
    ["Android OEM battery suppression", "Courier background tracking may pause on some devices.", "Foreground service, stopOnTerminate=false, startOnBoot=true, heartbeat interval, and field-validation recommendation."],
    ["Network dead zones", "Live tracking pulses may fail during travel.", "SQLite tracking_outbox, batch-sync endpoint, retry schedule, dead-letter handling, and connectivity audit."],
    ["Route deviation false positives", "Traffic detours may be misclassified as misconduct.", "MOVING_POSITIVELY state distinguishes productive off-corridor movement from stationary deviation."],
    ["Background worker failure", "Matching, notification, tracking validation, or cleanup may silently stop.", "job_heartbeats table, scheduler heartbeats, and heartbeat_watchdog STUCK marking."],
    ["Privacy leakage", "Client could infer courier movement beyond delivery need.", "Raw coordinates stored privately; client receives progress, ETA, freshness, and integrity label only."],
], [1.65, 2.3, 2.2])

heading(doc, "4.13  Conclusion", 2)
para(doc, "This chapter has presented the system design and analysis of DropCity as an implemented distributed logistics platform. The design satisfies the DCE analysis-and-design expectations by defining hardware and software specifications, functional and non-functional requirements, component structure, API decomposition, data flow, physical deployment, database design, interface design, security design, and verification tests.")
para(doc, "The central technical contribution is the integration of corridor-constrained matching, privacy-preserving tracking, offline courier telemetry, evidence-backed handoff, and heartbeat-monitored backend workers into one coherent system. Chapter 5 can therefore proceed from design to implementation and results, presenting the concrete SQL procedures, Express endpoints, screenshots, pilot simulation evidence, and measured behaviour of the deployed prototype.")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(OUT))
print(OUT)
