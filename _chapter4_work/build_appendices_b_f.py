from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(r"C:\Users\HUAWEI\projects\courier\DropCity_Chapter4_SystemDesignAnalysis.docx")
OUT = Path(r"C:\Users\HUAWEI\projects\courier\DropCity_Appendices_B_F.docx")

SLATE = "2F4F4F"
TRANSIT = "008080"
LIGHT = "F8F9FA"
CODE_FILL = "F3F6F6"
BORDER = "94A3B8"


def clear_doc(doc):
    body = doc._body._element
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(table, color=BORDER):
    tblPr = table._tbl.tblPr
    tbl_borders = tblPr.first_child_found_in("w:tblBorders")
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tblPr.append(tbl_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = tbl_borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tbl_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_widths(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx >= len(row.cells):
                continue
            cell = row.cells[idx]
            cell.width = Inches(width)
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(int(width * 1440)))
            tcW.set(qn("w:type"), "dxa")


def cell_text(cell, text, bold=False, color=None, size=8.8, font="Arial"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = font
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def para(doc, text="", style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.font.name = "Arial"
    return p


def heading(doc, text, level):
    return doc.add_paragraph(text, style=f"Heading {level}")


def table_caption(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.runs[0]
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    r.bold = True


def caption(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.runs[0]
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    r.italic = True


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(headers):
        shade_cell(t.rows[0].cells[i], SLATE)
        cell_text(t.rows[0].cells[i], h, bold=True, color="FFFFFF", size=9)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            shade_cell(cells[i], LIGHT if len(t.rows) % 2 == 0 else "FFFFFF")
            cell_text(cells[i], value)
    set_widths(t, widths)
    borders(t)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def bullet(doc, text):
    try:
        p = doc.add_paragraph(style="List Bullet")
    except Exception:
        p = doc.add_paragraph(style="List Paragraph")
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.name = "Arial"


def screenshot_placeholder(doc, title, instructions, fig_no):
    t = doc.add_table(rows=2, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_widths(t, [6.3])
    shade_cell(t.rows[0].cells[0], SLATE)
    cell_text(t.rows[0].cells[0], f"SCREENSHOT PLACEHOLDER: {title}", bold=True, color="FFFFFF", size=9.4)
    shade_cell(t.rows[1].cells[0], LIGHT)
    cell_text(t.rows[1].cells[0], f"Placement Instructions: {instructions}", size=8.8)
    borders(t)
    caption(doc, f"Figure B.{fig_no}: {title}.")


def code_block(doc, title, source, code):
    table_caption(doc, title)
    t = doc.add_table(rows=2, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_widths(t, [6.3])
    shade_cell(t.rows[0].cells[0], SLATE)
    cell_text(t.rows[0].cells[0], source, bold=True, color="FFFFFF", size=8.6)
    shade_cell(t.rows[1].cells[0], CODE_FILL)
    cell = t.rows[1].cells[0]
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for idx, line in enumerate(code.rstrip().splitlines()):
        if idx:
            p.add_run("\n")
        run = p.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(7.4)
    borders(t)


def read_lines(path, start, end):
    lines = Path(path).read_text(encoding="utf-8").splitlines()
    return "\n".join(lines[start - 1:end])


doc = Document(str(BASE))
clear_doc(doc)
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)

for style_name in ["Normal", "List Paragraph"]:
    try:
        st = doc.styles[style_name]
        st.font.name = "Arial"
        st.font.size = Pt(11)
        st.paragraph_format.line_spacing = 1.08
        st.paragraph_format.space_after = Pt(8)
    except Exception:
        pass

heading(doc, "APPENDIX B: USER MANUAL", 1)
heading(doc, "B.1  Introduction", 2)
para(doc, "This user manual is designed for DropCity clients, couriers, recipients, and administrators. It explains how to operate the platform without requiring specialised technical knowledge of PostGIS, Flutter, Firebase, or backend services.")
para(doc, "DropCity is a resilient peer-to-peer logistics platform that matches parcels to couriers already travelling along compatible corridors. It uses corridor-constrained matching, secure pickup and dropoff handovers, privacy-preserving progress tracking, and administrative monitoring to support urban delivery in connectivity-constrained environments.")

heading(doc, "B.2  System Overview", 2)
table_caption(doc, "Table B.1: User roles and primary responsibilities")
table(doc, ["Role", "Main Responsibility", "Primary Screens or Console Areas"], [
    ["Client", "Create parcel requests, track parcel progress, complete sender-side handoff actions, and rate courier service.", "Client Home, Deliveries, Delivery Creation Flow, Parcel Status, Sender/Recipient Handoff screens."],
    ["Courier", "Register vehicle details, declare corridors, accept parcel requests, complete pickup/dropoff handovers, and keep tracking active.", "Courier Home, Routes, Parcels, Pickup Mode, Pickup, Dropoff, Settings."],
    ["Recipient", "Receive parcel status notifications and issue or confirm dropoff verification where the in-app recipient flow is used.", "Recipient Handoff and Parcel Status screens."],
    ["Administrator", "Monitor platform health, approve vehicles/couriers, inspect disputes, resolve alerts, and review spatial analytics.", "Admin Monitoring, Health, Alerts, Disputes, Couriers, Vehicles, Scheduler, Spatial Analytics."],
], [1.2, 3.0, 2.25])

heading(doc, "B.3  System Requirements", 2)
table_caption(doc, "Table B.2: Minimum operational requirements")
table(doc, ["Area", "Requirement", "Notes"], [
    ["Mobile device", "Android/iOS device capable of running the Flutter app build.", "Courier devices must support reliable GPS and background location permissions."],
    ["Network", "Mobile data or Wi-Fi connection.", "Temporary disconnection is supported through the courier outbox, but account setup and sync require reconnection."],
    ["Location", "GPS/location services enabled.", "Required for corridor declaration, pickup/dropoff proximity gates, and tracking updates."],
    ["Camera", "Camera access enabled.", "Required for pickup/dropoff photographic proof."],
    ["Notifications", "Push notification permission recommended.", "Used for courier acceptance, PIN readiness, ETA/status changes, alerts, and handoff events."],
    ["Admin browser", "Modern browser with JavaScript enabled.", "Used to access the Next.js admin console."],
], [1.25, 2.75, 2.2])

heading(doc, "B.4  Getting Started", 2)
heading(doc, "B.4.1  Account Registration and Verification", 3)
para(doc, "New users should select the correct account type before completing onboarding. Clients complete personal profile details, while couriers complete additional operational verification steps. Identity documentation must be uploaded clearly so the administrator can verify that the courier profile is legitimate before live delivery operations.")
bullet(doc, "Open the DropCity app and select the correct role: Client or Courier.")
bullet(doc, "Enter email, password, and personal details in the multi-step signup screens.")
bullet(doc, "Upload required identity documentation where prompted.")
bullet(doc, "Wait for account or vehicle approval if the selected role requires administrative verification.")
screenshot_placeholder(doc, "Client and Courier Signup Flow", "Place this screenshot immediately after Section B.4.1. Capture the multi-step signup screens showing email entry, personal information, role selection, and document upload prompts. Use clean demo data and ensure no real ID number, phone number, or private email address is visible.", 1)

heading(doc, "B.4.2  Vehicle Registration for Couriers", 3)
para(doc, "Couriers must register the vehicle they intend to use for corridor-based delivery. The registration should include licence information, plate details, and supporting vehicle images where the app requests them. The administrator reviews vehicle records before the courier is treated as fully operational.")
bullet(doc, "Open the courier onboarding or settings workflow.")
bullet(doc, "Enter licence number, vehicle make/model, and plate number.")
bullet(doc, "Upload licence and vehicle images using the camera or gallery option.")
bullet(doc, "Check the approval status from the courier dashboard or admin confirmation message.")
screenshot_placeholder(doc, "Courier Vehicle Registration and Approval Status", "Place this screenshot after the vehicle registration instructions. Capture the courier vehicle details form and a second view showing pending or approved status. The screenshot should make the licence/plate workflow visible without exposing a real plate number.", 2)

heading(doc, "B.5  Courier Operations", 2)
heading(doc, "B.5.1  Declaring Corridors", 3)
para(doc, "A corridor is the courier's declared route path. DropCity uses this route to determine whether a parcel can be carried without requiring the courier to make an economically unreasonable detour.")
bullet(doc, "Open the Courier app and go to Routes.")
bullet(doc, "Select the route declaration option and choose the starting and ending points on the map.")
bullet(doc, "Confirm the route geometry, route timing, and any route details requested by the app.")
bullet(doc, "Submit the route so it becomes available for matching.")
screenshot_placeholder(doc, "Courier Route Declaration Map", "Place this screenshot after Section B.5.1. Capture the map route declaration screen with a clear start marker, end marker, route polyline, and confirmation button. If possible, use a Harare demo route such as CBD to University of Zimbabwe.", 3)

heading(doc, "B.5.2  The Active Workflow", 3)
para(doc, "After declaring a route, the courier must activate the journey before accepting and executing live parcel work. When tracking is active, the app uses the background geolocation engine and local outbox so movement updates can continue even when connectivity is unstable.")
bullet(doc, "Open the route details or courier dashboard screen.")
bullet(doc, "Start or activate the route before collecting assigned parcels.")
bullet(doc, "Keep location permission enabled and do not force-stop the courier application during active deliveries.")
bullet(doc, "Watch the tracking health indicators and outbox/dead-letter counters if they are shown on the dashboard.")
screenshot_placeholder(doc, "Courier Active Route and Tracking Health", "Place this screenshot after Section B.5.2. Capture the courier dashboard or route details view showing an active route state, tracking indicator, and any outbox/health counters. The image should show that tracking is active without exposing private coordinates.", 4)

heading(doc, "B.5.3  Secure Handovers Using the Three-Way Handshake", 3)
para(doc, "The Three-Way Handshake protects custody transfer by requiring a PIN or OTP, physical proximity, and photographic proof. Pickup changes the parcel to IN_TRANSIT; dropoff changes it to COMPLETED.")
table_caption(doc, "Table B.3: Courier handover procedure")
table(doc, ["Step", "Courier Action", "Expected Result"], [
    ["1", "Open the assigned parcel and select Pickup.", "The app prompts for pickup PIN, GPS position, and photo proof."],
    ["2", "Meet the sender at the approved pickup or meeting point.", "The courier should be within the GPS gate before verification is accepted."],
    ["3", "Enter the pickup PIN and capture the pickup photo.", "A successful pickup moves the parcel to IN_TRANSIT."],
    ["4", "Travel toward the dropoff point while tracking remains active.", "The client sees progress and integrity status, not raw courier coordinates."],
    ["5", "At dropoff, enter the dropoff PIN or OTP and capture proof.", "A successful dropoff moves the parcel to COMPLETED."],
], [0.65, 3.0, 2.55])
screenshot_placeholder(doc, "Pickup and Dropoff Three-Way Handshake Screens", "Place this screenshot after Table B.3. Use a two-panel composite: left panel showing pickup PIN/photo/proximity prompt; right panel showing dropoff OTP/photo prompt. The screenshot should clearly show the camera proof area and verification button.", 5)

heading(doc, "B.6  Client Operations", 2)
heading(doc, "B.6.1  Creating a Parcel", 3)
para(doc, "Clients create parcels through a guided multi-step flow. The form should capture what is being delivered, its size or weight category, pickup location, dropoff location, recipient details, and any priority or fragility information.")
bullet(doc, "Open the Client app and select the delivery creation option.")
bullet(doc, "Enter parcel description, size, weight, priority, and fragility where applicable.")
bullet(doc, "Choose pickup and dropoff coordinates on the map or location picker.")
bullet(doc, "Review the delivery summary and submit the parcel request.")
screenshot_placeholder(doc, "Client Parcel Creation Flow", "Place this screenshot after Section B.6.1. Capture the delivery creation flow showing parcel details, map-based pickup/dropoff selection, and final confirmation summary. Use demo information only.", 6)

heading(doc, "B.6.2  Tracking Progress", 3)
para(doc, "The client status page intentionally avoids raw courier live tracking. Instead, it displays parcel status, progress percentage, ETA where available, tracking freshness, and an integrity label such as On planned route, Off-route but progressing, Potential route deviation, or Tracking nominal.")
table_caption(doc, "Table B.4: Tracking labels and user interpretation")
table(doc, ["Label", "Meaning", "Recommended User Response"], [
    ["On planned route", "The courier's pulse is within the expected corridor.", "Continue monitoring normally."],
    ["Off-route but progressing", "The courier is outside the route corridor but still moving toward the destination.", "Allow time for traffic detour recovery unless the status persists."],
    ["Potential route deviation", "The system has detected stationary or suspicious off-corridor behaviour.", "Check alerts or contact support/admin if the delivery appears unsafe."],
    ["Tracking pending", "No recent tracking pulse is available yet.", "Wait for the courier app to send or sync a pulse."],
], [1.7, 2.75, 1.7])
screenshot_placeholder(doc, "Client Parcel Status and Progress Checkpoints", "Place this screenshot after Table B.4. Capture the parcel status screen with progress percentage, ETA/freshness text, integrity label, and handoff timeline. Do not include raw courier coordinates.", 7)

heading(doc, "B.7  Admin Console", 2)
heading(doc, "B.7.1  Monitoring Dashboard", 3)
para(doc, "The monitoring dashboard is used to inspect platform health. Administrators should pay special attention to heartbeat status, stuck jobs, notification backlog, tracking observability, and recent error logs.")
bullet(doc, "Open the admin console and sign in with an administrator account.")
bullet(doc, "Review heartbeat cards for ACTIVE, STUCK, or inactive job states.")
bullet(doc, "Open scheduler or health views to identify jobs that have not reported within their expected interval.")
bullet(doc, "Review alerts and logs before restarting or investigating a backend worker.")
screenshot_placeholder(doc, "Admin Monitoring Dashboard with Heartbeat Alerts", "Place this screenshot after Section B.7.1. Capture the admin health or monitoring page showing heartbeat records, status labels, and a visible example of healthy versus stuck job presentation. If no job is stuck, annotate the screenshot with a callout placeholder for where STUCK would appear.", 8)

heading(doc, "B.7.2  Dispute Resolution", 3)
para(doc, "The dispute workspace helps administrators resolve claims by reviewing handoff evidence, GPS logs, status transitions, and user reports. The administrator should compare the photographic evidence and location/time sequence before deciding whether the courier, client, or recipient claim is supported.")
bullet(doc, "Open the Admin Disputes page.")
bullet(doc, "Select the relevant parcel or dispute record.")
bullet(doc, "Review pickup/dropoff photos, handshake events, and tracking logs.")
bullet(doc, "Record a resolution decision with clear notes.")
screenshot_placeholder(doc, "Admin Dispute Evidence Workspace", "Place this screenshot after Section B.7.2. Capture the dispute detail page showing parcel metadata, pickup/dropoff evidence, tracking log summary, and resolution controls. Blur faces, phone numbers, or private address text before final submission.", 9)

heading(doc, "B.8  Troubleshooting", 2)
table_caption(doc, "Table B.5: Common issues and recovery actions")
table(doc, ["Issue", "Likely Cause", "Recovery Action"], [
    ["Courier tracking not updating", "No network, location permission disabled, or background service paused.", "Enable location permission, restore network, reopen Courier app, and allow outbox sync."],
    ["Outbox count increasing", "Device is recording pulses but cannot reach backend.", "Continue delivery if safe; queued pulses should sync automatically once network returns."],
    ["Dead-letter count increasing", "Updates exceeded retry limit or were rejected.", "Report to admin/developer with parcel ID and approximate time."],
    ["Pickup PIN rejected", "Wrong PIN or repeated failed attempts caused lockout.", "Verify PIN with sender and wait for lockout window if triggered."],
    ["Outside GPS gate", "Courier is not within the accepted handoff area.", "Move closer to the pickup/dropoff or approved meeting point and retry."],
    ["Admin sees STUCK job", "Background worker missed the heartbeat threshold.", "Inspect logs, restart worker if needed, and confirm heartbeat returns to ACTIVE."],
], [1.65, 2.4, 2.15])
screenshot_placeholder(doc, "Courier Dead-Zone Recovery Indicators", "Place this screenshot at the end of Appendix B. Capture the courier dashboard or settings area showing outbox count, dead-letter count, last sync time, or tracking health message. The screenshot should explain how a courier recognises that offline pulses are queued rather than lost.", 10)

doc.add_page_break()
heading(doc, "APPENDIX F: SOURCE CODE SNIPPETS", 1)
para(doc, "This appendix presents selected source code snippets that demonstrate the core technical mechanisms behind DropCity's academic contribution. The snippets were selected from the current repository and focus on spatial matching, offline-first synchronization, heartbeat reliability monitoring, and Three-Way Handshake security.")

heading(doc, "F.1  Spatial Matching Logic (PostgreSQL/PostGIS)", 2)
code_block(doc, "Code Snippet F.1: Corridor-constrained parcel matching", "Source: backend/sql/002_matching.sql", read_lines("backend/sql/002_matching.sql", 1, 31))
para(doc, "Significance: This stored procedure is the geospatial core of DropCity. It uses ST_DWithin to require both parcel endpoints to lie inside the courier corridor and ST_LineLocatePoint to prove that pickup occurs before dropoff. This supports the zero-detour claim by filtering for parcels already compatible with the courier's declared journey rather than computing a new delivery trip.")

heading(doc, "F.2  Offline-First Synchronization (Flutter/SQLite)", 2)
para(doc, "The current implementation does not use a literal is_synced boolean flag. Instead, unsynced pulses remain in the tracking_outbox table until they are accepted by /tracking/batch-sync; successful or permanently skipped rows are deleted, while repeated failures are retried or moved to the dead-letter table.")
code_block(doc, "Code Snippet F.2: SQLite outbox schema and FIFO fetch", "Source: courier/lib/services/tracking_outbox.dart", read_lines("courier/lib/services/tracking_outbox.dart", 44, 106))
code_block(doc, "Code Snippet F.3: Reconnection flush and batch-sync processing", "Source: courier/lib/services/courier_tracking_service.dart", read_lines("courier/lib/services/courier_tracking_service.dart", 315, 394))
para(doc, "Significance: These snippets validate DropCity's resilient distributed architecture. The courier phone is treated as an edge node that can store ordered location pulses locally, wait for network recovery, replay pulses in FIFO order, and isolate permanently failed records in a dead-letter queue.")

heading(doc, "F.3  Watchdog Heartbeat Monitor (Node.js)", 2)
code_block(doc, "Code Snippet F.4: Background worker heartbeat watchdog", "Source: backend/src/jobs/heartbeat_watchdog.js", read_lines("backend/src/jobs/heartbeat_watchdog.js", 1, 42))
para(doc, "Significance: This function proves that backend reliability is observable rather than assumed. Each worker has an expected frequency; if the current time exceeds twice that interval since last_heartbeat_at, the watchdog marks the job as STUCK in job_heartbeats for administrator visibility.")

heading(doc, "F.4  Three-Way Handshake Hashing and Geofence Gate", 2)
code_block(doc, "Code Snippet F.5: SHA-256 PIN hashing and GPS gate setup", "Source: backend/src/routes/handshake.js", read_lines("backend/src/routes/handshake.js", 1, 24))
code_block(doc, "Code Snippet F.6: Pickup PIN and 50m proximity validation", "Source: backend/src/routes/handshake.js", read_lines("backend/src/routes/handshake.js", 285, 365))
para(doc, "Significance: These snippets validate the cryptographic and spatial security implementation. The system hashes PINs with SHA-256, checks repeated failed attempts, verifies the courier against a 50-metre GPS gate, records failed GPS attempts, and only transitions the parcel to IN_TRANSIT after successful PIN, proximity, and evidence checks.")

heading(doc, "F.5  Appendix Summary", 2)
table_caption(doc, "Table F.1: Academic significance of selected implementation snippets")
table(doc, ["Snippet", "Core Mechanism", "Academic Contribution"], [
    ["F.1", "PostGIS ST_DWithin and ST_LineLocatePoint corridor matching", "Demonstrates spatially constrained opportunistic logistics."],
    ["F.2-F.3", "SQLite outbox, FIFO fetch, batch sync, retry and dead-letter handling", "Demonstrates offline-first mobile edge resilience."],
    ["F.4", "Heartbeat watchdog and STUCK transition", "Demonstrates self-monitoring process reliability."],
    ["F.5-F.6", "SHA-256 PIN verification and 50m geofence validation", "Demonstrates evidence-backed secure physical custody transfer."],
], [1.0, 2.65, 2.55])

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(OUT))
print(OUT)
